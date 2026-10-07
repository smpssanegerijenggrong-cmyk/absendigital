"""End-to-end HTTP/API tests, SQLite temporer terisolasi, tanpa internet."""
import io
import os
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

os.environ['SECRET_KEY']='TEST-only-sanjara-key-at-least-32-characters-2026'
os.environ['ADMIN_USERNAME']='admin'
os.environ['ADMIN_PASSWORD']='TestPasswordSanjara_2026'

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook,load_workbook
from app import main as m,db,backup,services

WIB=ZoneInfo('Asia/Jakarta')
TODAY=datetime(2026,10,7,7,12,30,tzinfo=WIB)

@pytest.fixture
def system(tmp_path,monkeypatch):
    dbpath=tmp_path/'database.sqlite3'
    letters=tmp_path/'letters';letters.mkdir()
    monkeypatch.setattr(db,'DB_PATH',dbpath)
    monkeypatch.setattr(db,'LETTERS_DIR',letters)
    monkeypatch.setattr(backup,'DB_PATH',dbpath)
    monkeypatch.setattr(backup,'LETTERS_DIR',letters)
    monkeypatch.setattr(m,'LETTERS_DIR',letters)
    monkeypatch.setattr(m,'now_wib',lambda:TODAY)
    monkeypatch.setattr(services,'now_wib',lambda:TODAY)
    db.init_db();m.bootstrap_admin()
    with TestClient(m.app) as client:
        yield client

def csrf(client,url='/login'):
    text=client.get(url).text
    match=re.search(r'name="_csrf" value="([^"]+)"',text)
    assert match,(url,text[:250])
    return match.group(1)

def login(client):
    key=csrf(client)
    r=client.post('/login',data={'_csrf':key,'username':'admin','password':os.environ['ADMIN_PASSWORD']},follow_redirects=False)
    assert r.status_code==303,r.text
    return key

def add_student(client,name='Ani',nisn='1234567890',nipd='11',klass='VII A'):
    r=client.post('/students/add',data={'_csrf':csrf(client,'/students'),'name':name,'gender':'P',
        'nisn':nisn,'nipd':nipd,'class_name':klass},follow_redirects=True)
    assert r.status_code==200
    with db.connect() as con:
        row=con.execute('SELECT * FROM students WHERE name=?',(name,)).fetchone()
        assert row,r.text[:600]
        return dict(row)

def gps():
    with db.connect() as con:
        db.set_setting(con,'latitude','-8.10');db.set_setting(con,'longitude','113.10')
        db.set_setting(con,'radius_m','150');db.set_setting(con,'workdays','0,1,2,3,4,5,6')

def scan(client,student,lat=-8.10,lon=113.10,accuracy=12):
    return client.post('/api/scan',json={'token':'SANJARA:'+student['qr_token'],
        'latitude':lat,'longitude':lon,'accuracy':accuracy},headers={'X-CSRF-Token':csrf(client,'/scan')})

def test_login_csrf_and_role(system):
    c=system
    assert c.get('/').status_code==200 # redirect followed to login
    assert 'Selamat datang kembali' in c.get('/login').text
    assert c.post('/login',data={'username':'admin','password':'irrelevant'}).status_code==403
    login(c)
    assert c.get('/').status_code==200
    assert 'SANJARA HADIR' in c.get('/').text
    assert c.get('/health').json()['status']=='ok'
    assert c.post('/students/add',data={'name':'Bad'}).status_code==403
    assert c.get('/reports').status_code==200
    assert c.get('/settings').status_code==200


def test_classes_students_and_cards(system):
    c=system;login(c)
    student=add_student(c)
    for url in ['/','/classes','/students','/cards','/students/'+str(student['id'])+'/card','/leave','/reports','/settings','/users','/scan']:
        r=c.get(url)
        assert r.status_code==200,(url,r.status_code,r.text[:160])
        assert 'SANJARA' in r.text
    qr=c.get(f'/students/{student["id"]}/qr.png')
    assert qr.status_code==200 and qr.headers['content-type']=='image/png'
    assert qr.content.startswith(b'\x89PNG')
    c.post(f'/students/{student["id"]}/toggle',data={'_csrf':csrf(c,'/students')})
    with db.connect() as con: assert con.execute('SELECT active FROM students WHERE id=?',(student['id'],)).fetchone()['active']==0


def test_cutoff_exact_and_milliseconds():
    assert services.late_status(datetime(2026,10,7,7,0,tzinfo=WIB),'07:00:00')==('HADIR',0)
    assert services.late_status(datetime(2026,10,7,7,0,1,tzinfo=WIB),'07:00:00')==('TERLAMBAT',1)
    assert services.late_status(datetime(2026,10,7,7,12,30,tzinfo=WIB),'07:00:00')==('TERLAMBAT',750)
    assert services.late_status(datetime(2026,10,7,7,0,0,1000,tzinfo=WIB),'07:00:00')==('TERLAMBAT',1)


def test_scan_gps_cutoff_duplicate_and_absence(system):
    c=system;login(c);gps()
    student=add_student(c)
    bad=scan(c,student,lat=-8.12)
    assert bad.status_code==403
    assert scan(c,student,accuracy=120).status_code==422
    good=scan(c,student)
    assert good.status_code==200,good.text
    output=good.json()
    assert output['student']=='Ani' and output['status']=='TERLAMBAT'
    assert output['time']=='07:12:30' and output['late_seconds']==750
    dup=scan(c,student)
    assert dup.status_code==409
    with db.connect() as con:
        records=con.execute('SELECT * FROM attendance WHERE student_id=?',(student['id'],)).fetchall()
        assert len(records)==1 and records[0]['source']=='QR'
    with db.connect() as con:
        db.set_setting(con,'closing','07:10:00')
    another=add_student(c,name='Budi',nisn='2222',nipd='22')
    assert scan(c,another).status_code==409
    with db.connect() as con:
        opts=db.settings(con)
        assert services.visible_status(con,another['id'],'2026-10-07',opts)== 'ALPA'


def test_import_students_and_classes(system):
    c=system;login(c)
    key=csrf(c,'/classes')
    stream=io.BytesIO(b'KELAS\nVII A\nVIII A\nVII A\n')
    r=c.post('/classes/import',data={'_csrf':key},files={'file':('kelas.csv',stream,'text/csv')},follow_redirects=True)
    assert r.status_code==200
    with db.connect() as con:assert con.execute('SELECT COUNT(*) FROM classes').fetchone()[0]==2
    wb=Workbook();ws=wb.active;ws.append(['NIPD','NISN','NAMA','JENIS KELAMIN','KELAS'])
    ws.append(['100','991','Aulia','Perempuan','VII A']);ws.append(['101','992','Bima','Laki-laki','IX A'])
    stream=io.BytesIO();wb.save(stream);stream.seek(0)
    r=c.post('/students/import',data={'_csrf':csrf(c,'/students')},files={'file':('murid.xlsx',stream,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')},follow_redirects=True)
    assert r.status_code==200 and '2 siswa baru' in r.text
    with db.connect() as con:assert con.execute('SELECT COUNT(*) FROM students').fetchone()[0]==2
    assert c.get('/students/template').status_code==200
    assert c.get('/classes/template').status_code==200
    assert 'IX A' in c.get('/classes').text


def test_leave_approval_and_excel(system):
    c=system;login(c);gps()
    student=add_student(c,name='Dewi',nisn='5555',nipd='55')
    form={'_csrf':csrf(c,'/leave'),'student_id':student['id'],'day':'2026-10-07',
          'kind':'IJIN','reason':'Lainnya','custom_reason':'Menghadiri acara keluarga'}
    r=c.post('/leave/add',data=form,files={'letter':('surat.pdf',b'%PDF-1.4\n1 0 obj\nendobj\n%%EOF','application/pdf')},follow_redirects=True)
    assert r.status_code==200 and 'menunggu persetujuan' in r.text
    with db.connect() as con:
        req=con.execute('SELECT * FROM leave_requests').fetchone()
        assert req['reason']=='Menghadiri acara keluarga'
        assert req['state']=='PENDING'
    r=c.post(f'/leave/{req["id"]}/review',data={'_csrf':csrf(c,'/leave'),'action':'APPROVED'},follow_redirects=True)
    assert r.status_code==200 and 'disetujui' in r.text
    with db.connect() as con:
        row=con.execute('SELECT status,source FROM attendance WHERE student_id=?',(student['id'],)).fetchone()
        assert dict(row)=={'status':'IJIN','source':'SURAT'}
    response=c.get(f'/leave/{req["id"]}/letter')
    assert response.status_code==200 and response.content.startswith(b'%PDF')
    report=c.get('/reports/export.xlsx?month=2026-10&day=2026-10-07')
    assert report.status_code==200
    wb=load_workbook(io.BytesIO(report.content))
    assert wb.sheetnames==['REKAP','DETAIL WAKTU','HARIAN','PETUNJUK','KALENDER LIBUR']
    assert wb['REKAP']['J1'].value=='JUMLAH'
    assert wb['REKAP']['K2'].value=='IJIN'
    assert wb['REKAP']['H3'].value==1
    assert wb['DETAIL WAKTU']['G2'].value=='IJIN'
    backup_response=c.get('/backup/download')
    assert backup_response.status_code==200
    with zipfile.ZipFile(io.BytesIO(backup_response.content)) as zipf:
        assert 'sanjara.sqlite3' in zipf.namelist()
        assert any(name.startswith('letters/') for name in zipf.namelist())


def test_leave_reject_and_holiday(system):
    c=system;login(c);gps()
    student=add_student(c,'Hana','6666','66')
    form={'_csrf':csrf(c,'/leave'),'student_id':student['id'],'day':'2026-10-07',
          'kind':'SAKIT','reason':'Berobat / sakit'}
    c.post('/leave/add',data=form,files={'letter':('surat.png',b'\x89PNG\r\n\x1a\nanything','image/png')})
    with db.connect() as con: item=con.execute('SELECT id FROM leave_requests').fetchone()
    c.post(f'/leave/{item["id"]}/review',data={'_csrf':csrf(c,'/leave'),'action':'REJECTED'})
    with db.connect() as con:
        assert con.execute('SELECT COUNT(*) FROM attendance').fetchone()[0]==0
    c.post('/settings/holiday',data={'_csrf':csrf(c,'/settings'),'day':'2026-10-07','description':'Libur sekolah'})
    with db.connect() as con: assert services.visible_status(con,student['id'],'2026-10-07',db.settings(con))=='LIBUR'
    assert scan(c,student).status_code==409


def test_petugas_role_and_no_default_authorization(system):
    c=system;login(c)
    r=c.post('/users/add',data={'_csrf':csrf(c,'/users'),'username':'petugas1','display_name':'Petugas Satu',
                                'password':'PetugasPassword2026','role':'PETUGAS'},follow_redirects=True)
    assert r.status_code==200
    logoutcsrf=csrf(c,'/users')
    c.post('/logout',data={'_csrf':logoutcsrf})
    c.get('/login')
    r=c.post('/login',data={'_csrf':csrf(c),'username':'petugas1','password':'PetugasPassword2026'},follow_redirects=False)
    assert r.status_code==303
    assert c.get('/scan').status_code==200
    assert c.get('/students').status_code==200
    assert c.get('/settings').status_code==403
    assert c.get('/users').status_code==403
    assert c.post('/students/add',data={'_csrf':csrf(c,'/students'),'name':'X'}).status_code==403


def test_drive_config_default(system):
    c=system;login(c)
    assert c.get('/drive/connect',follow_redirects=False).status_code==303
    assert 'backup' in c.get('/settings').text.lower()
    assert c.post('/backup/drive',data={'_csrf':csrf(c,'/settings')},follow_redirects=True).status_code==200


def test_student_edit_qr_rotation_and_class_rename(system):
    c=system;login(c)
    student=add_student(c,'Putri','7777','77')
    editpage=c.get(f'/students/{student["id"]}/edit')
    assert editpage.status_code==200 and 'Edit identitas siswa' in editpage.text
    with db.connect() as con: klass_id=con.execute('SELECT id FROM classes').fetchone()['id']
    c.post(f'/classes/{klass_id}/rename',data={'_csrf':csrf(c,'/classes'),'name':'VII B'})
    assert 'VII B' in c.get('/classes').text
    r=c.post(f'/students/{student["id"]}/edit',data={'_csrf':csrf(c,f'/students/{student["id"]}/edit'),
              'name':'Putri Amelia','nipd':'77','nisn':'7777','gender':'P','class_name':'VII B'},follow_redirects=True)
    assert 'Putri Amelia' in r.text
    c.post(f'/students/{student["id"]}/new-qr',data={'_csrf':csrf(c,f'/students/{student["id"]}/edit')})
    with db.connect() as con:
        record=con.execute('SELECT qr_token,name FROM students WHERE id=?',(student['id'],)).fetchone()
        assert record['qr_token']!=student['qr_token']
        assert record['name']=='Putri Amelia'


def test_password_change_profile_and_failed_attempts(system):
    c=system;login(c)
    assert c.get('/profile').status_code==200
    r=c.post('/profile/password',data={'_csrf':csrf(c,'/profile'),
              'current_password':os.environ['ADMIN_PASSWORD'],'new_password':'ChangedPassword2026',
              'confirm_password':'ChangedPassword2026'},follow_redirects=True)
    assert r.status_code==200 and 'berhasil diganti' in r.text
    assert 'Masuk sebagai administrator' in r.text
    r=c.post('/login',data={'_csrf':csrf(c),'username':'admin','password':'ChangedPassword2026'},
              follow_redirects=False)
    assert r.status_code==303
    c.post('/logout',data={'_csrf':csrf(c,'/profile')})
    for _ in range(10):
        c.post('/login',data={'_csrf':csrf(c),'username':'admin','password':'WrongPassword'})
    failed=c.post('/login',data={'_csrf':csrf(c),'username':'admin','password':'ChangedPassword2026'},follow_redirects=True)
    assert 'Terlalu banyak percobaan' in failed.text


def test_drive_oauth_flow_mocked(system,monkeypatch):
    c=system;login(c)
    monkeypatch.setenv('GDRIVE_CLIENT_ID','client-test')
    monkeypatch.setenv('GDRIVE_CLIENT_SECRET','secret-test')
    monkeypatch.setenv('GDRIVE_REDIRECT_URI','http://testserver/drive/callback')
    monkeypatch.setattr(m.drive,'exchange_code',lambda code,uri:'fake-google-refresh-token')
    connected=c.get('/drive/connect',follow_redirects=False)
    assert connected.status_code in (302,307)
    from urllib.parse import urlparse,parse_qs
    params=parse_qs(urlparse(connected.headers['location']).query)
    assert 'state' in params and 'client_id' in params
    r=c.get('/drive/callback',params={'state':params['state'][0],'code':'example'},follow_redirects=True)
    assert r.status_code==200 and 'berhasil dihubungkan' in r.text
    with db.connect() as con:assert db.settings(con)['drive_refresh_token_enc'].startswith('gAAAA')
    def fake_upload(dbconn,archive,filename):
        assert archive.getbuffer().nbytes>1000
        assert filename.endswith('.zip')
        return {'id':'file-123','name':filename}
    monkeypatch.setattr(m.drive,'upload_backup',fake_upload)
    r=c.post('/backup/drive',data={'_csrf':csrf(c,'/settings')},follow_redirects=True)
    assert r.status_code==200 and 'Backup berhasil diunggah' in r.text
    r=c.post('/drive/disconnect',data={'_csrf':csrf(c,'/settings')},follow_redirects=True)
    assert r.status_code==200 and 'diputus' in r.text


def test_restore_backup_offline(system,tmp_path,monkeypatch):
    c=system;login(c);student=add_student(c,'Restorable','9090','90')
    backupdata=c.get('/backup/download').content
    archive=tmp_path/'backup.zip';archive.write_bytes(backupdata)
    c.post(f'/students/{student["id"]}/toggle',data={'_csrf':csrf(c,'/students')})
    with db.connect() as con:assert con.execute('SELECT active FROM students WHERE id=?',(student['id'],)).fetchone()['active']==0
    import scripts.restore_backup as restoremod
    monkeypatch.setattr(restoremod,'DB_PATH',db.DB_PATH)
    monkeypatch.setattr(restoremod,'LETTERS_DIR',db.LETTERS_DIR)
    monkeypatch.setattr(restoremod,'DATA_DIR',tmp_path)
    restoremod.restore(archive) # unit test simulated maintenance, no running real HTTP server
    with db.connect() as con:assert con.execute('SELECT active FROM students WHERE id=?',(student['id'],)).fetchone()['active']==1
