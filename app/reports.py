"""Ekspor XLSX kompatibel rekap sekolah + detail jam absensi."""
import io
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
from .services import attendance_report,parse_month,visible_status,now_wib
from .db import settings,connect

TEAL='087E8B'; NAVY='13283F'; LIGHT='E9F4F5'

def clean(value):
    value = '' if value is None else str(value)
    return "'"+value if value.startswith(('=','+','-','@')) else value

def decorate(ws,widths,heading_rows=(1,),freeze='A2'):
    ws.freeze_panes=freeze
    ws.sheet_view.showGridLines=False
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    for rowid in heading_rows:
        for cell in ws[rowid]:
            cell.fill=PatternFill('solid',fgColor=NAVY if rowid==1 else TEAL)
            cell.font=Font(color='FFFFFF',bold=True,size=10)
            cell.alignment=Alignment(vertical='center',horizontal='center',wrap_text=True)
        ws.row_dimensions[rowid].height=28
    ws.auto_filter.ref=ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage=True
    ws.print_options.horizontalCentered=True
    ws.page_setup.orientation='landscape'
    ws.page_setup.paperSize=ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth=1
    for row in ws.iter_rows(min_row=max(heading_rows)+1):
        for cell in row:
            cell.alignment=Alignment(vertical='center')
            if cell.row%2==0: cell.fill=PatternFill('solid',fgColor='F1F7F8')

def make_report(db,month,day,klass=''):
    records=attendance_report(db,month,day,klass)
    options=settings(db)
    wb=Workbook(); ws=wb.active;ws.title='REKAP'
    ws.append(['No','NIPD','NISN','NAMA','JENIS KELAMIN','KELAS','SAKIT','IJIN','ALPA','JUMLAH',None,None])
    ws.merge_cells('J1:L1')
    ws.append(['','','','','','','','','','SAKIT','IJIN','ALPA'])
    for i,r in enumerate(records,1):
        s=r['student']; tot=r['totals'];status=r['status']
        ws.append([i,clean(s['nipd']),clean(s['nisn']),clean(s['name']),s['gender'],clean(s['class_name']),
                   int(status=='SAKIT'),int(status=='IJIN'),int(status=='ALPA'),tot['SAKIT'],tot['IJIN'],tot['ALPA']])
    decorate(ws,[7,16,18,35,19,16,13,13,13,15,15,15],(1,2),'A3')
    ws.auto_filter.ref=f'A2:L{max(2,ws.max_row)}'
    detail=wb.create_sheet('DETAIL WAKTU')
    detail.append(['TANGGAL','NIPD','NISN','NAMA SISWA','KELAS','JAM ABSEN WIB','STATUS','TERLAMBAT (MENIT)','SUMBER','PETUGAS'])
    rows=db.execute('''SELECT a.*,s.nipd,s.nisn,s.name,c.name AS class_name,u.display_name AS petugas
        FROM attendance a JOIN students s ON s.id=a.student_id
        JOIN classes c ON c.id=s.class_id LEFT JOIN users u ON u.id=a.operator_id
        WHERE a.day BETWEEN ? AND ? AND (?='' OR c.name=?) ORDER BY a.day,c.name,s.name''',
        (month+'-01',month+'-31',klass,klass)).fetchall()
    for r in rows:
        detail.append([r['day'],clean(r['nipd']),clean(r['nisn']),clean(r['name']),clean(r['class_name']),
                       r['scanned_at'][11:19] if r['scanned_at'] else '',r['status'],
                       round((r['late_seconds'] or 0)/60,2),r['source'],clean(r['petugas'])])
    decorate(detail,[17,16,18,35,17,19,18,23,16,23])
    daily=wb.create_sheet('HARIAN')
    daily.append(['TANGGAL','NIPD','NISN','NAMA','KELAS','STATUS','JAM ABSEN WIB'])
    for d in parse_month(month):
        for r in records:
            s=r['student']
            if d < s['joined_on'] or d>now_wib().date().isoformat(): continue
            state=visible_status(db,s['id'],d,options,s['joined_on'])
            mark=db.execute('SELECT scanned_at FROM attendance WHERE student_id=? AND day=?',(s['id'],d)).fetchone()
            daily.append([d,clean(s['nipd']),clean(s['nisn']),clean(s['name']),clean(s['class_name']),state,
                          mark['scanned_at'][11:19] if mark and mark['scanned_at'] else ''])
    decorate(daily,[17,18,18,35,17,21,19])
    legend=wb.create_sheet('PETUNJUK')
    for row in [
       ['SANJARA HADIR — REKAP ABSENSI',options['school_name']],
       ['Tanggal rekap',day],['Bulan rekap',month],['Kelas',klass or 'Semua kelas'],
       ['Kolom G–I','SAKIT/IJIN/ALPA khusus tanggal pilihan (0 atau 1)'],
       ['Kolom J–L','Jumlah SAKIT/IJIN/ALPA selama bulan pilihan'],
       ['HADIR','Scan sampai 07:00:00 WIB (bisa diubah admin)'],
       ['TERLAMBAT','Scan di atas jam batas, tetap dianggap hadir'],
       ['ALPA','Hari sekolah yang telah ditutup, tanpa absensi/izin disetujui'],
       ['DETAIL WAKTU','Waktu setiap scan dan nama petugas'],
       ['HARIAN','Satu baris per siswa per hari untuk pemantauan'],
    ]: legend.append(row)
    legend.column_dimensions['A'].width=33;legend.column_dimensions['B'].width=80
    legend['A1'].font=Font(bold=True,color=TEAL,size=14)
    calendar=wb.create_sheet('KALENDER LIBUR')
    calendar.append(['TANGGAL','KETERANGAN'])
    for row in db.execute('SELECT day,description FROM holidays WHERE day BETWEEN ? AND ? ORDER BY day',
                          (month+'-01',month+'-31')): calendar.append([row['day'],clean(row['description'])])
    decorate(calendar,[20,70])
    buf=io.BytesIO();wb.save(buf);buf.seek(0)
    return buf
